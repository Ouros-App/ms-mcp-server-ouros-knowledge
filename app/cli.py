import argparse
import hashlib
import json
import re
import uuid
from collections.abc import Iterable
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from app.core.config import settings
from app.services.knowledge import get_qdrant_client, get_vector_store, qdrant_status

DEFAULT_DOCS_DIR = Path("docs")
MANIFEST_NAME = ".qdrant-manifest.json"
SUPPORTED_EXTENSIONS = {".csv", ".docx", ".html", ".htm", ".json", ".md", ".pdf", ".txt"}


def split_text(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    if chunk_size <= 0:
        raise ValueError("chunk-size deve ser maior que zero")
    if not 0 <= chunk_overlap < chunk_size:
        raise ValueError("chunk-overlap deve estar entre zero e chunk-size - 1")

    text = re.sub(r"[ \t]+\n", "\n", text).strip()
    chunks: list[str] = []
    start = 0
    while start < len(text):
        hard_end = min(start + chunk_size, len(text))
        end = hard_end
        if hard_end < len(text):
            break_at = max(text.rfind("\n", start, hard_end), text.rfind(" ", start, hard_end))
            if break_at > start + chunk_size // 2:
                end = break_at

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(end - chunk_overlap, start + 1)
    return chunks


def discover_files(paths: Iterable[Path]) -> list[Path]:
    files: set[Path] = set()
    for path in paths:
        path = path.expanduser()
        if not path.exists():
            raise FileNotFoundError(f"caminho não encontrado: {path}")
        if path.is_file():
            if path.name == MANIFEST_NAME:
                continue
            if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                raise ValueError(f"formato não suportado: {path.suffix or path.name}")
            files.add(path.resolve())
            continue
        files.update(
            candidate.resolve()
            for candidate in path.rglob("*")
            if candidate.is_file()
            and candidate.name != MANIFEST_NAME
            and candidate.suffix.lower() in SUPPORTED_EXTENSIONS
        )
    return sorted(files)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_manifest(path: Path) -> dict:
    if not path.exists():
        return {"version": 1, "files": {}}
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise RuntimeError(f"manifesto inválido: {path}") from error
    if not isinstance(manifest, dict) or not isinstance(manifest.get("files", {}), dict):
        raise TypeError(f"manifesto inválido: {path}")
    manifest.setdefault("files", {})
    return manifest


def save_manifest(path: Path, manifest: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def manifest_key(path: Path, docs_dir: Path) -> str:
    try:
        return path.relative_to(docs_dir).as_posix()
    except ValueError:
        return path.as_posix()


def extract_file(path: Path) -> list[Document]:
    metadata = {
        "source": str(path),
        "file_name": path.name,
        "file_type": path.suffix.lower().lstrip("."),
    }
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        documents = []
        for page_number, page in enumerate(PdfReader(str(path)).pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                documents.append(Document(page_content=text, metadata={**metadata, "page": page_number}))
        return documents

    if suffix == ".docx":
        document = DocxDocument(str(path))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        blocks.extend(
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        )
        text = "\n".join(blocks)
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")

    return [Document(page_content=text, metadata=metadata)] if text.strip() else []


def prepare_documents(path: Path, chunk_size: int, chunk_overlap: int) -> list[Document]:
    chunks: list[Document] = []
    for source_document in extract_file(path):
        for chunk_index, text in enumerate(
            split_text(source_document.page_content, chunk_size, chunk_overlap)
        ):
            chunks.append(
                Document(
                    page_content=text,
                    metadata={**source_document.metadata, "chunk_index": chunk_index},
                )
            )
    return chunks


def document_id(document: Document) -> str:
    source = document.metadata["source"]
    page = document.metadata.get("page", "")
    chunk_index = document.metadata["chunk_index"]
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source}:{page}:{chunk_index}"))


def ingest(
    paths: list[Path],
    chunk_size: int,
    chunk_overlap: int,
    batch_size: int,
    dry_run: bool,
    manifest_path: Path,
) -> int:
    files = discover_files(paths)
    if not files:
        print("Nenhum arquivo suportado encontrado.")
        return 0

    docs_dir = DEFAULT_DOCS_DIR.resolve()
    manifest = load_manifest(manifest_path)
    manifest["collection"] = settings.QDRANT_COLLECTION_NAME
    manifest["embedding_model"] = settings.NVIDIA_EMBEDDING_MODEL
    pending: list[tuple[Path, str, list[Document], dict]] = []
    for path in files:
        key = manifest_key(path, docs_dir)
        digest = file_hash(path)
        previous = manifest["files"].get(key, {})
        changed = (
            previous.get("sha256") != digest
            or previous.get("collection") != settings.QDRANT_COLLECTION_NAME
            or previous.get("embedding_model") != settings.NVIDIA_EMBEDDING_MODEL
        )
        if not changed:
            print(f"IGNORADO sem alteração: {key}")
            continue

        documents = prepare_documents(path, chunk_size, chunk_overlap)
        ids = [document_id(document) for document in documents]
        record = {
            "sha256": digest,
            "collection": settings.QDRANT_COLLECTION_NAME,
            "embedding_model": settings.NVIDIA_EMBEDDING_MODEL,
            "chunk_count": len(documents),
            "ids": ids,
        }
        pending.append((path, key, documents, record))
        print(f"ALTERADO: {key} -> {len(documents)} chunks")

    if not pending:
        print("Nenhuma alteração para enviar.")
        return 0

    store = None
    if not dry_run:
        status = qdrant_status()
        if not status["exists"]:
            raise RuntimeError(
                f"coleção Qdrant não encontrada: {status['collection']} ({status['error']})"
            )
        store = get_vector_store()
    client = get_qdrant_client() if not dry_run else None

    total = 0
    for path, key, documents, record in pending:
        if dry_run:
            continue
        for start in range(0, len(documents), batch_size):
            batch = documents[start : start + batch_size]
            store.add_documents(batch, ids=[document_id(document) for document in batch])
            total += len(batch)

        previous_ids = set(manifest["files"].get(key, {}).get("ids", []))
        stale_ids = sorted(previous_ids - set(record["ids"]))
        if stale_ids:
            client.delete(
                collection_name=settings.QDRANT_COLLECTION_NAME,
                points_selector=stale_ids,
                wait=True,
            )
        manifest["files"][key] = record
        save_manifest(manifest_path, manifest)

    if dry_run:
        print("Dry-run concluído; nada foi enviado ao Qdrant.")
    else:
        print(f"Upload concluído: {total} chunks enviados para {store.collection_name}.")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Embeddar e enviar arquivos para o Qdrant.")
    subparsers = parser.add_subparsers(dest="command", required=True)
    ingest_parser = subparsers.add_parser("ingest", help="processa arquivos e envia os chunks")
    ingest_parser.add_argument(
        "paths", nargs="*", type=Path, help="arquivos ou diretórios (padrão: ./docs)"
    )
    ingest_parser.add_argument("--chunk-size", type=int, default=1000)
    ingest_parser.add_argument("--chunk-overlap", type=int, default=150)
    ingest_parser.add_argument("--batch-size", type=int, default=32)
    ingest_parser.add_argument("--manifest", type=Path, default=DEFAULT_DOCS_DIR / MANIFEST_NAME)
    ingest_parser.add_argument("--dry-run", action="store_true")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if args.command == "ingest":
        if args.batch_size <= 0:
            raise SystemExit("batch-size deve ser maior que zero")
        try:
            return ingest(
                args.paths or [DEFAULT_DOCS_DIR],
                args.chunk_size,
                args.chunk_overlap,
                args.batch_size,
                args.dry_run,
                args.manifest,
            )
        except (FileNotFoundError, TypeError, ValueError, RuntimeError) as error:
            raise SystemExit(f"Erro: {error}") from error
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
