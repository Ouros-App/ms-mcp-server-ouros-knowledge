import base64
import binascii
import json
import zipfile
from io import BytesIO
from typing import Any
from urllib.request import Request, urlopen

from pypdf import PdfReader

from app.core.config import settings

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 100
MAX_XLSX_ROWS = 10_000
MAX_XLSX_CELLS = 100_000
ALLOWED_TYPES = {"application/pdf", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
DOCUMENT_TOO_LARGE = "documento convertido excede o limite de importação"


def _xlsx_sheet_lines(sheet: Any):
    """Yield bounded Markdown lines for one worksheet."""
    header: list[str] | None = None
    cell_count = 0
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number > MAX_XLSX_ROWS:
            raise ValueError("XLSX excede o limite de linhas")
        row = ["" if value is None else str(value) for value in values]
        cell_count += len(row)
        if cell_count > MAX_XLSX_CELLS:
            raise ValueError("XLSX excede o limite de células")
        if not any(cell.strip() for cell in row):
            continue
        if header is None:
            header = row
            yield f"## Planilha: {sheet.title}\n"
            yield "| " + " | ".join(row) + " |\n"
            yield "| " + " | ".join(["---"] * len(row)) + " |\n"
            continue
        normalized = row[:len(header)] + [""] * max(0, len(header) - len(row))
        yield "| " + " | ".join(normalized) + " |\n"


def _xlsx_to_markdown(content: bytes) -> str:
    """Convert workbook cells into a bounded Markdown document."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("suporte XLSX não está instalado") from error
    with zipfile.ZipFile(BytesIO(content)) as archive:
        uncompressed_size = sum(item.file_size for item in archive.infolist())
        compressed_size = max(sum(item.compress_size for item in archive.infolist()), 1)
    if uncompressed_size > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError("XLSX excede o limite descompactado")
    if uncompressed_size / compressed_size > MAX_XLSX_COMPRESSION_RATIO:
        raise ValueError("taxa de compressão do XLSX excede o limite")

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections: list[str] = []
    total_chars = 0
    try:
        for sheet in workbook.worksheets:
            for line in _xlsx_sheet_lines(sheet):
                total_chars += len(line)
                if total_chars > settings.IMPORT_MARKDOWN_MAX_CHARS:
                    raise ValueError(DOCUMENT_TOO_LARGE)
                sections.append(line)
    finally:
        workbook.close()
    return "\n".join(sections)


def _pdf_to_markdown(content: bytes) -> str:
    """Extract text from a text-based PDF into Markdown."""
    reader = PdfReader(BytesIO(content))
    return "\n\n".join(
        f"## Página {index}\n\n{page.extract_text() or ''}"
        for index, page in enumerate(reader.pages, start=1)
    )


def file_to_markdown(content_type: str, encoded_file: str) -> str:
    """Decode and convert one supported PDF/XLSX upload to Markdown."""
    if content_type not in ALLOWED_TYPES:
        raise ValueError("tipo de arquivo não suportado; use PDF ou XLSX")
    try:
        content = base64.b64decode(encoded_file, validate=True)
    except binascii.Error as error:
        raise ValueError("arquivo deve ser base64 válido") from error
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("arquivo vazio ou maior que 8 MiB")
    if content_type == "application/pdf":
        if not content.startswith(b"%PDF"):
            raise ValueError("assinatura de PDF inválida")
        markdown = _pdf_to_markdown(content)
    else:
        if not zipfile.is_zipfile(BytesIO(content)):
            raise ValueError("assinatura de XLSX inválida")
        markdown = _xlsx_to_markdown(content)
    if not markdown.strip():
        raise ValueError("não foi possível extrair texto do arquivo")
    if len(markdown) > settings.IMPORT_MARKDOWN_MAX_CHARS:
        raise ValueError(DOCUMENT_TOO_LARGE)
    return markdown


def extract_resource_records(markdown: str, source_name: str) -> dict[str, Any]:
    """Ask an NVIDIA NIM to extract canonical records without performing writes."""
    if settings.NVIDIA_API_KEY is None or not settings.NVIDIA_NIM_URL:
        raise RuntimeError("NVIDIA NIM não está configurado")
    prompt = (
        "Extraia somente registros históricos de água e energia deste documento. "
        "Responda JSON com a chave records, usando exatamente os campos "
        "resource_type, farm_id, registration_date, start_hydrometer, "
        "end_hydrometer, energy_consumption, source_row e confidence. "
        "Não invente valores; use null quando ausente. Documento: " + source_name
    )
    payload = {"model": settings.NVIDIA_NIM_MODEL, "temperature": 0,
               "response_format": {"type": "json_object"},
               "messages": [{"role": "system", "content": prompt},
                            {"role": "user", "content": markdown}]}
    request = Request(settings.NVIDIA_NIM_URL, data=json.dumps(payload).encode(),
                      headers={
                          "Authorization": (
                              "Bearer "
                              + settings.NVIDIA_API_KEY.get_secret_value()
                          ),
                          "Content-Type": "application/json",
                      },
                      method="POST")
    try:
        with urlopen(request, timeout=settings.NVIDIA_NIM_TIMEOUT) as response:
            body = json.loads(response.read())
        content = body["choices"][0]["message"]["content"]
        result = json.loads(content) if isinstance(content, str) else content
        records = result.get("records")
        if not isinstance(records, list):
            raise TypeError("NIM não retornou uma lista de records")
        return {"source_name": source_name, "records": records, "preview": True}
    except (KeyError, TypeError, ValueError) as error:
        raise RuntimeError("resposta inválida do NVIDIA NIM") from error
